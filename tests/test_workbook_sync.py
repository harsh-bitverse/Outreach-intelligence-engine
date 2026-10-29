"""
Unit and Integration Tests for Dynamic Workbook Synchronization and Atomic Updates.
"""

import os
import shutil
import tempfile
from pathlib import Path
import openpyxl
import pytest
from src.database.connection import get_db, initialize_database
from src.database.repository import OutreachRepository, compute_durable_professor_id
from src.database.workbook_sync import DynamicWorkbookSync, WorkbookLockError
from src.database.models import Professor


@pytest.fixture
def temp_db_and_workbook():
    """Create a temporary SQLite DB and sample Excel workbook."""
    temp_dir = Path(tempfile.mkdtemp())
    db_path = temp_dir / "test_outreach.db"
    wb_path = temp_dir / "test_professors.xlsx"

    # Initialize DB
    initialize_database(db_path)

    # Create synthetic test workbook with dynamic headers and custom formulas
    wb = openpyxl.Workbook()
    ws_prof = wb.active
    ws_prof.title = "Professor Database"

    headers = [
        "QS Rank", "Processing Sequence", "University Name", "Country",
        "Professor Full Name", "Academic Title / Position", "Department / Faculty",
        "Official Institutional Email", "Email Verification Status",
        "Email Verification Source URL", "Official Faculty Profile URL",
        "Research Areas / Keywords", "Specific Research Overlap With My Interests",
        "Relevant Skill or Project From My CV", "Potential Software / Computational Contribution",
        "Research Evidence URL", "Date Verified", "Notes / Unresolved Issues"
    ]
    ws_prof.append(headers)

    # Row 1: Valid professor
    ws_prof.append([
        1, 1, "Massachusetts Institute of Technology (MIT)", "United States",
        "Abigail Bodner", "Assistant Professor", "EECS",
        "bodner@mit.edu", "VERIFIED", "https://mit.edu/source",
        "https://mit.edu/bodner", "Climate modeling, ML", "Ocean dynamics and SciML",
        "Neural Process Twin", "Surrogate modeling in PyTorch",
        "https://mit.edu/evidence", "2026-10-25", "Screened and verified"
    ])

    # Row 2: Suppressed professor (Kai-Kit Wong)
    ws_prof.append([
        9, 9, "University College London (UCL)", "United Kingdom",
        "Kai-Kit Wong", "Chair Professor", "Electronic and Electrical Engineering",
        "k.wong@ucl.ac.uk", "VERIFIED", "https://ucl.ac.uk/source",
        "https://ucl.ac.uk/wong", "Fluid antenna systems, wireless", "Wireless comms",
        "Neural Process Twin", "Optimization",
        "https://ucl.ac.uk/evidence", "2026-10-25", "Screened"
    ])

    # Row 3: Suppressed institutional domain (IIT Delhi)
    ws_prof.append([
        150, 50, "Indian Institute of Technology Delhi (IITD)", "India",
        "Test Professor IITD", "Professor", "Chemical Engineering",
        "faculty@chemical.iitd.ac.in", "VERIFIED", "https://iitd.ac.in",
        "https://iitd.ac.in/prof", "Transport phenomena", "Fluid mechanics",
        "Neural Process Twin", "Numerical simulation",
        "https://iitd.ac.in", "2026-10-25", "Internal"
    ])

    wb.save(str(wb_path))
    wb.close()

    yield db_path, wb_path, temp_dir

    shutil.rmtree(temp_dir, ignore_errors=True)


def test_durable_professor_id_consistency():
    """Verify durable professor ID is deterministic and insensitive to whitespace/case."""
    id1 = compute_durable_professor_id("MIT", "Abigail Bodner")
    id2 = compute_durable_professor_id("  mit  ", "abigail bodner  ")
    assert id1 == id2
    assert id1.startswith("prof_")
    assert len(id1) == 21  # 'prof_' + 16 chars


def test_dynamic_schema_import(temp_db_and_workbook):
    """Verify workbook import reads records dynamically and suppresses pre-configured targets."""
    db_path, wb_path, _ = temp_db_and_workbook
    sync = DynamicWorkbookSync(wb_path)

    with get_db(db_path) as conn:
        repo = OutreachRepository(conn)
        stats = sync.import_to_database(repo)

        assert stats["professors_imported"] == 3

        # Verify Abigail Bodner is imported with status NEW
        abigail_id = compute_durable_professor_id("Massachusetts Institute of Technology (MIT)", "Abigail Bodner")
        abigail = repo.get_professor(abigail_id)
        assert abigail is not None
        assert abigail["email"] == "bodner@mit.edu"
        assert abigail["outreach_status"] == "NEW"

        # Verify Kai-Kit Wong is automatically flagged as SUPPRESSED
        wong_id = compute_durable_professor_id("University College London (UCL)", "Kai-Kit Wong")
        wong = repo.get_professor(wong_id)
        assert wong is not None
        assert wong["outreach_status"] == "SUPPRESSED"

        # Verify IIT Delhi professor is automatically flagged as SUPPRESSED
        iitd_id = compute_durable_professor_id("Indian Institute of Technology Delhi (IITD)", "Test Professor IITD")
        iitd_prof = repo.get_professor(iitd_id)
        assert iitd_prof is not None
        assert iitd_prof["outreach_status"] == "SUPPRESSED"


def test_atomic_sync_back_and_backup_creation(temp_db_and_workbook):
    """Verify sync back creates timestamped backup and updates outreach status non-destructively."""
    db_path, wb_path, temp_dir = temp_db_and_workbook
    sync = DynamicWorkbookSync(wb_path)

    with get_db(db_path) as conn:
        repo = OutreachRepository(conn)
        sync.import_to_database(repo)

        # Update Abigail Bodner status in SQLite
        abigail_id = compute_durable_professor_id("Massachusetts Institute of Technology (MIT)", "Abigail Bodner")
        repo.update_professor_status(abigail_id, "APPROVED")

        # Mock time blocks data
        sample_blocks = [{
            "block_id": "BLOCK_09_US_EAST",
            "region": "US Eastern",
            "ist_time": "6:30 PM IST",
            "local_time": "8:00 AM - 9:00 AM Local",
            "countries": "United States",
            "university_count": 1,
            "faculty_count": 1,
            "notes": "Test block"
        }]

        # Write back to workbook atomically
        sync.sync_back_to_workbook(repo, sample_blocks)

        # 1. Verify backup file was created
        backups = list(temp_dir.glob("*.bak"))
        assert len(backups) >= 1

        # 2. Verify Excel content
        wb_updated = openpyxl.load_workbook(str(wb_path))
        assert "Time Blocks and Distribution" in wb_updated.sheetnames

        ws_prof = wb_updated["Professor Database"]
        # Find Outreach Status column
        headers = [cell.value for cell in ws_prof[1]]
        assert "Outreach Status" in headers
        status_idx = headers.index("Outreach Status") + 1

        # Row 2 (Abigail) should be APPROVED
        assert ws_prof.cell(row=2, column=status_idx).value == "APPROVED"
        # Row 3 (Kai-Kit Wong) should be SUPPRESSED
        assert ws_prof.cell(row=3, column=status_idx).value == "SUPPRESSED"
        wb_updated.close()


def test_workbook_lock_prevention(temp_db_and_workbook):
    """Verify lock file prevents concurrent overlapping writes."""
    _, wb_path, _ = temp_db_and_workbook
    sync1 = DynamicWorkbookSync(wb_path)
    sync2 = DynamicWorkbookSync(wb_path)

    sync1.acquire_lock()
    try:
        with pytest.raises(WorkbookLockError):
            sync2.acquire_lock(timeout_seconds=1)
    finally:
        sync1.release_lock()
